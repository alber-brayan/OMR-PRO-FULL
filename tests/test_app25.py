"""Pruebas del flujo local, aislamiento de datos, revisión y exportación."""
import io
import json
from pathlib import Path
import tempfile
import types
import unittest
from unittest.mock import patch
import zipfile

import numpy as np
from PIL import Image

from app import create_app


def image_bytes(color='white'):
    image = Image.new('RGB', (1400, 990), color)
    data = io.BytesIO()
    image.save(data, 'JPEG')
    return data.getvalue()


def fake_read(path):
    return {'name': 'MUÑOZ PEREZ ANA', 'school': 'SAN RAMON', 'grade': 2,
            'answers': list('ABCDEABCDEABCDE'), 'issues': [], 'needs_review': False,
            'quality': {'sharpness': 100.0}, 'warped_image': np.full((990, 1400, 3), 255, np.uint8)}


def fake_grade(answers, key, scoring=None):
    if '?' in answers:
        raise ValueError('Respuesta dudosa')
    scoring = scoring or {'buenas': 10, 'malas': -1, 'blanco': 0}
    correct = sum(a == b for a, b in zip(answers, key))
    blank = answers.count('-')
    incorrect = 15-correct-blank
    return {'correct': correct, 'incorrect': incorrect, 'blank': blank,
            'double': answers.count('X'), 'score': correct*scoring['buenas']+incorrect*scoring['malas']+blank*scoring['blanco'],
            'results': [{'question': i+1, 'answer': a, 'expected': b, 'status': 'correct' if a == b else 'blank' if a == '-' else 'incorrect'} for i, (a, b) in enumerate(zip(answers, key))]}


class App25Tests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.mock_engine = types.SimpleNamespace(
            read_sheet=fake_read, grade_answers=fake_grade,
            annotate_reading=lambda reading, grading=None: reading['warped_image'],
            build_sheet_pdf=lambda **kw: (b'%PDF-1.4\n', 'ficha_omr25_a4.pdf'))
        self.fake_excel = types.SimpleNamespace(export_report=lambda records, **kw: b'PK-TEST-EXPORT')
        self.modules = patch.dict('sys.modules', {'omr25': self.mock_engine, 'excel25': self.fake_excel})
        self.modules.start()
        self.app = create_app(self.tmp.name, testing=True)
        self.client = self.app.test_client()

    def tearDown(self):
        self.modules.stop()
        self.tmp.cleanup()

    def key(self, course='Matematica', grade=2, value='ABCDEABCDEABCDE'):
        return self.client.post('/api/keys', json={'course': course, 'grade': grade, 'answers': value})

    def scan(self, data=None):
        return self.client.post('/api/scan', data={'course': 'Matematica', 'files': (io.BytesIO(data or image_bytes()), 'ficha.jpg')}, content_type='multipart/form-data')

    def confirm(self, record_id, **overrides):
        return self.client.post(f'/api/records/{record_id}/confirm', json={'reviewed': True, **overrides})

    def test_pages_and_download_disposition(self):
        self.assertEqual(self.client.get('/').status_code, 200)
        self.assertEqual(self.client.get('/sheet').status_code, 200)
        pdf = self.client.get('/download_sheet_pdf?inline=1')
        self.assertTrue(pdf.headers['Content-Disposition'].startswith('inline'))
        self.assertEqual(self.client.get('/sheet?curso=Otro').status_code, 400)
        self.assertEqual(self.client.get('/no-existe').status_code, 404)

    def test_review_required_missing_key_and_manual_correction(self):
        result = self.scan()
        self.assertEqual(result.status_code, 200)
        record = result.json['records'][0]
        self.assertFalse(record['confirmed'])
        self.assertTrue(record['needs_review'])
        self.assertEqual(self.confirm(record['id']).status_code, 409)
        self.assertEqual(self.key().status_code, 200)
        self.assertEqual(self.client.post(f"/api/records/{record['id']}/confirm", json={}).status_code, 400)
        confirmed = self.confirm(record['id'], name='MUÑOZ PEREZ ANA MARIA', school='COLEGIO SAN RAMON', order='023', time='10:35')
        self.assertEqual(confirmed.status_code, 200, confirmed.json)
        self.assertEqual(confirmed.json['correct'], 15)
        self.assertEqual(confirmed.json['score'], 150)
        self.assertEqual(confirmed.json['order'], '023')
        self.assertEqual(confirmed.json['time'], '10:35')
        self.assertTrue(confirmed.json['confirmed'])
        self.assertEqual(self.client.get(f"/api/records/{record['id']}/image").status_code, 200)

    def test_uncertainty_and_invalid_manual_fields_cannot_confirm(self):
        self.key()
        record = self.scan().json['records'][0]
        for payload in ({'answers': ['?']*15}, {'grade': 7}, {'name': ''}, {'school': ''}, {'name': 'MU?OZ ANA'}, {'school': 'SAN?JOSE'}, {'order': '../etc'}, {'time': '25:10'}, {'answers': 'A'*14}):
            with self.subTest(payload=payload):
                response = self.confirm(record['id'], **payload)
                self.assertEqual(response.status_code, 400, response.json)
        self.assertFalse(self.client.get('/api/records').json['records'][0]['confirmed'])

    def test_duplicate_images_do_not_duplicate_records(self):
        self.key()
        first = self.scan().json['records'][0]
        second = self.scan().json['records'][0]
        self.assertEqual(first['id'], second['id'])
        self.assertTrue(second['duplicate'])
        self.assertEqual(len(self.client.get('/api/records').json['records']), 1)

    def test_key_and_scoring_changes_invalidate_results(self):
        self.key()
        record = self.scan().json['records'][0]
        self.assertEqual(self.confirm(record['id']).status_code, 200)
        self.key(value='A'*15)
        self.assertFalse(self.client.get('/api/records').json['records'][0]['confirmed'])
        self.assertEqual(self.confirm(record['id']).status_code, 200)
        response = self.client.post('/api/settings', json={'scoring': {'buenas': 5, 'malas': 0, 'blanco': 0}})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(self.client.get('/api/records').json['records'][0]['confirmed'])

    def test_exports_include_only_confirmed_current_course(self):
        self.assertEqual(self.client.get('/api/export/excel').status_code, 409)
        self.key()
        record = self.scan().json['records'][0]
        self.assertEqual(self.client.get('/api/export/pdf').status_code, 409)
        self.confirm(record['id'])
        report = self.client.get('/api/export/excel')
        self.assertEqual(report.status_code, 200)
        self.assertEqual(report.data, b'PK-TEST-EXPORT')
        pdf = self.client.get('/api/export/pdf')
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.data.startswith(b'%PDF'))
        self.assertEqual(self.client.get('/api/export/excel?course=Comunicacion').status_code, 409)
        self.assertEqual(self.client.delete(f"/api/records/{record['id']}").status_code, 200)
        self.assertEqual(self.client.get('/api/export/excel').status_code, 409)

    def test_bad_images_and_zip_paths_preserve_previous_records(self):
        record = self.scan().json['records'][0]
        response = self.scan(b'not-an-image')
        self.assertEqual(response.status_code, 422)
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, 'w') as zf:
            zf.writestr('../bad.jpg', image_bytes())
        response = self.client.post('/api/scan', data={'course':'Matematica','files':(io.BytesIO(archive.getvalue()),'malicioso.zip')}, content_type='multipart/form-data')
        self.assertEqual(response.status_code, 422)
        self.assertEqual(self.client.get('/api/records').json['records'][0]['id'], record['id'])
        self.assertEqual(self.client.get('/api/records/../../foo/image').status_code, 404)

    def test_zip_accepts_folder_layout_and_reports_individual_errors(self):
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, 'w') as zf:
            zf.writestr('Sec 2/Matematica/001.jpg', image_bytes())
            zf.writestr('Sec 2/Matematica/mala.jpg', b'bad')
            zf.writestr('manifest.json', '{}')
        response = self.client.post('/api/scan', data={'course':'Matematica','files':(io.BytesIO(archive.getvalue()),'pack.zip')}, content_type='multipart/form-data')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['count'], 1)
        self.assertEqual(len(response.json['errors']), 1)

    def test_folder_paths_with_same_basename_keep_both_photos(self):
        response = self.client.post('/api/scan', data={'course':'Matematica', 'files':[
            (io.BytesIO(image_bytes('white')), 'Lote/Grado_1/ficha.jpg'),
            (io.BytesIO(image_bytes('gray')), 'Lote/Grado_6/ficha.jpg'),
        ]}, content_type='multipart/form-data')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['count'], 2)
        self.assertEqual({r['source_path'] for r in response.json['records']},
                         {'Lote/Grado_1/ficha.jpg','Lote/Grado_6/ficha.jpg'})
        self.assertEqual(len({r['id'] for r in response.json['records']}), 2)

    def test_bulk_requires_review_and_excludes_ambiguous_records(self):
        self.key()
        clean = self.scan().json['records'][0]
        original = self.mock_engine.read_sheet
        def ambiguous(path):
            reading = fake_read(path)
            reading.update(name='MU?OZ ANA', issues=['Una letra es dudosa.'], needs_review=True)
            return reading
        self.mock_engine.read_sheet = ambiguous
        uncertain = self.scan(image_bytes('gray')).json['records'][0]
        self.mock_engine.read_sheet = original
        self.assertEqual(self.client.post('/api/records/confirm-clear', json={'course':'Matematica'}).status_code, 400)
        result = self.client.post('/api/records/confirm-clear', json={'course':'Matematica','reviewed':True})
        self.assertEqual(result.status_code, 200, result.json)
        self.assertEqual(result.json['count'], 1)
        self.assertEqual(result.json['excluded'], 1)
        records = {r['id']:r for r in self.client.get('/api/records').json['records']}
        self.assertTrue(records[clean['id']]['confirmed'])
        self.assertFalse(records[uncertain['id']]['confirmed'])

    def test_logo_ratio_and_deletion(self):
        response = self.client.post('/api/logo/left', data={'file':(io.BytesIO(image_bytes()),'logo.jpg')}, content_type='multipart/form-data')
        self.assertEqual(response.status_code, 200)
        png = self.client.get('/api/logo/left/image')
        with Image.open(io.BytesIO(png.data)) as image:
            self.assertEqual(image.size, (1500, 450))
        self.assertEqual(self.client.delete('/api/logo/left').status_code, 200)
        self.assertEqual(self.client.get('/api/logo/left/image').status_code, 404)
        self.assertEqual(self.client.post('/api/logo/other').status_code, 404)

    def test_keys_and_records_are_separate_by_course(self):
        self.key('Matematica', 2)
        self.assertIsNone(self.client.get('/api/keys?course=Comunicacion').json['keys']['2'])
        self.scan()
        self.assertEqual(self.client.get('/api/records?course=Comunicacion').json['records'], [])
        self.assertEqual(len(self.client.get('/api/records?course=Matematica').json['records']), 1)


class LiveIntegration25Tests(unittest.TestCase):
    def test_marked_pdf_to_review_to_excel_and_pdf(self):
        import cv2
        import pymupdf
        from tests.test_omr25_engine import render_sheet, desk
        from openpyxl import load_workbook
        answers = list('ABCDEABCDEABCDE')
        with tempfile.TemporaryDirectory() as directory:
            client = create_app(directory, testing=True).test_client()
            self.assertEqual(client.post('/api/keys', json={'course':'Matematica','grade':3,'answers':answers}).status_code, 200)
            source = desk(render_sheet(answers=answers), rotation=1, perspective=True, shadow=True)
            _, encoded = cv2.imencode('.png', source)
            read = client.post('/api/scan', data={'course':'Matematica','files':(io.BytesIO(encoded.tobytes()),'foto_prueba.png')}, content_type='multipart/form-data')
            self.assertEqual(read.status_code, 200, read.json)
            record = read.json['records'][0]
            self.assertEqual((record['name'],record['school'],record['grade']), ('MUÑOZ ANA','SAN JOSE',3))
            self.assertEqual(record['answers'], answers)
            self.assertEqual(record['issues'], [])
            confirmed = client.post(f"/api/records/{record['id']}/confirm", json={'reviewed':True,'order':'025','time':'10:42'})
            self.assertEqual(confirmed.status_code, 200, confirmed.json)
            self.assertEqual((confirmed.json['correct'],confirmed.json['incorrect'],confirmed.json['blank'],confirmed.json['score']), (15,0,0,150))
            excel = client.get('/api/export/excel?course=Matematica')
            self.assertEqual(excel.status_code, 200)
            book = load_workbook(io.BytesIO(excel.data), data_only=True)
            self.assertEqual(book.sheetnames, [f'Grado {g}' for g in range(1,7)])
            self.assertTrue(any(cell.value=='MUÑOZ ANA' for row in book['Grado 3'] for cell in row))
            book.close()
            pdf = client.get('/api/export/pdf?course=Matematica')
            self.assertEqual(pdf.status_code, 200)
            with pymupdf.open(stream=pdf.data,filetype='pdf') as doc:
                self.assertEqual(len(doc),1)
                self.assertGreater(doc[0].rect.width, doc[0].rect.height)


if __name__ == '__main__':
    unittest.main()
