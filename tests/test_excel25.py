"""Pruebas de exportación sobre el archivo real de diseño de la editorial."""
from copy import deepcopy
import hashlib
import io
from pathlib import Path
import tempfile
import unittest
from xml.dom import minidom
import zipfile

from openpyxl import load_workbook

from excel25 import DEFAULT_TEMPLATE, export_report


def record(**changes):
    result = {"id": "demo", "name": "ALUMNO DE PRUEBA", "school": "COLEGIO DEMO", "grade": 1,
              "answers": ["A"]*15, "correct": 15, "incorrect": 0, "blank": 0,
              "double": 0, "score": 150, "order": "001", "time": "08:30", "confirmed": True}
    result.update(changes)
    return result


def workbook(data, cached=False):
    return load_workbook(io.BytesIO(data), data_only=cached)


class Excel25Tests(unittest.TestCase):
    def test_six_sheets_counts_values_and_original_template(self):
        before = hashlib.sha256(DEFAULT_TEMPLATE.read_bytes()).hexdigest()
        data = export_report([record(grade=n, name=f"DEMO GRADO {n}") for n in range(1,7)])
        wb = workbook(data, cached=True)
        self.assertEqual(wb.sheetnames,[f"Grado {n}" for n in range(1,7)])
        for grade, sheet in enumerate(wb,1):
            self.assertEqual(sheet["B3"].value,f"DEMO GRADO {grade}")
            self.assertEqual(sheet["D2"].value,"Nº ORDEN")
            self.assertEqual(sheet["D3"].value,"001")
            self.assertEqual(sheet["E3"].value,15)
            self.assertEqual(sheet["F3"].value,0)
            self.assertEqual(sheet["G3"].value,0)
            self.assertEqual(sheet["H3"].value,150)
            self.assertEqual(sheet["A3"].value,1)
            self.assertIsNone(sheet["B4"].value)
            self.assertIsNone(sheet["H4"].value)
            self.assertEqual(len(sheet._images),1)
        self.assertEqual(hashlib.sha256(DEFAULT_TEMPLATE.read_bytes()).hexdigest(),before)
        self.assertEqual(before,"e4a39fdd8cc7a17af77a1a842757cd5c859d9e6fca9fe8c0c04234ac4a444687")

    def test_exact_style_dimensions_image_and_native_shapes_are_preserved(self):
        data = export_report([record()])
        source = load_workbook(DEFAULT_TEMPLATE)
        target = workbook(data)
        for sheet in target:
            self.assertEqual(str(sheet.merged_cells),str(source.active.merged_cells))
            for col in "ABCDEFGHI":
                self.assertEqual(sheet.column_dimensions[col].width,source.active.column_dimensions[col].width)
            for row in (1,2,3,5,6,12,13,72):
                self.assertEqual(sheet.row_dimensions[row].height,source.active.row_dimensions[row].height)
                for col in "ABCDEFGHI":
                    self.assertEqual(sheet[f"{col}{row}"]._style,source.active[f"{col}{row}"]._style)
        with zipfile.ZipFile(DEFAULT_TEMPLATE) as original, zipfile.ZipFile(io.BytesIO(data)) as output:
            self.assertEqual(output.read("xl/styles.xml"),original.read("xl/styles.xml"))
            self.assertEqual(output.read("xl/media/image1.jpg"),original.read("xl/media/image1.jpg"))
            for grade in range(1,7):
                xml=output.read(f"xl/drawings/drawing{grade}.xml").decode()
                self.assertIn(f"{grade}º GRADO",xml)
                self.assertIn(f"MATEMÁTICA - GRADO {grade}",xml)
                self.assertNotIn("HUANCAPI",xml)
                self.assertIn("OMR grado",xml)
            self.assertNotIn("xl/sharedStrings.xml",output.namelist())
            self.assertNotIn("xl/calcChain.xml",output.namelist())
            self.assertNotIn(b">TEST<",b"".join(output.read(name) for name in output.namelist() if name.endswith(".xml")))

    def test_only_confirmed_records_are_exported_and_empty_cells_are_not_zero(self):
        data=export_report([record(confirmed=False),record(grade=2)])
        wb=workbook(data,True)
        self.assertIsNone(wb["Grado 1"]["B3"].value)
        self.assertIsNone(wb["Grado 1"]["H3"].value)
        self.assertEqual(wb["Grado 2"]["B3"].value,"ALUMNO DE PRUEBA")
        self.assertIsNone(wb["Grado 6"]["E3"].value)

    def test_more_than_seventy_students_extend_styles_formulas_and_print_range(self):
        records=[record(name=f"DEMO {i:03d}",id=str(i),grade=2) for i in range(85)]
        data=export_report(records)
        sheet=workbook(data,True)["Grado 2"]
        self.assertEqual(sheet["B87"].value,"DEMO 084")
        self.assertEqual(sheet["E87"].value,15)
        self.assertIn("$I$87",str(sheet.print_area))
        self.assertEqual(sheet["B87"]._style,sheet["B72"]._style)
        self.assertEqual(sheet.row_dimensions[87].height,30)
        self.assertIn("E87",workbook(data)["Grado 2"]["H87"].value)

    def test_ranking_score_time_ties_and_blank_time(self):
        records=[record(name="SIN HORA",time=""), record(name="TARDE",time="10:00"),
                 record(name="PRIMERO",time="08:30"),record(name="EMPATADO",time="08:30"),
                 record(name="MENOR",correct=14,incorrect=1,score=139,time="00:00")]
        data=export_report(records)
        sheet=workbook(data,True)["Grado 1"]
        self.assertEqual([sheet[f"B{r}"].value for r in range(3,8)],
                         ["EMPATADO","PRIMERO","TARDE","SIN HORA","MENOR"])
        self.assertEqual([sheet[f"A{r}"].value for r in range(3,8)],[1,1,3,4,5])
        self.assertIsNone(sheet["I6"].value)
        self.assertEqual(sheet["I7"].value.hour,0)
        self.assertIn("RANK",workbook(data)["Grado 1"]["A3"].value)

    def test_formulas_and_cached_scores_support_custom_scoring(self):
        r=record(correct=5,incorrect=5,blank=5,answers=["A"]*10+["-"]*5,score=8.5)
        scoring={"buenas":2,"malas":-.5,"blanco":.2}
        data=export_report([r],scoring=scoring)
        self.assertEqual(workbook(data,True)["Grado 1"]["H3"].value,8.5)
        self.assertEqual(workbook(data)["Grado 1"]["H3"].value,"=E3*(2)+F3*(-0.5)+G3*(0.2)")

    def test_names_schools_order_are_literal_text_against_formula_injection(self):
        r=record(name='=HYPERLINK("http://example.invalid","X")',school="+SUM(A1:A2)",order="@DDE()")
        sheet=workbook(export_report([r]))["Grado 1"]
        for col,key in (("B","name"),("C","school"),("D","order")):
            self.assertEqual(sheet[f"{col}3"].value,r[key])
            self.assertEqual(sheet[f"{col}3"].data_type,"s")

    def test_incomplete_or_inconsistent_confirmed_records_fail_explicitly(self):
        changes=[{"grade":0},{"name":""},{"school":""},{"score":None},
                 {"correct":None},{"correct":14},{"score":1},{"score":float("nan")},
                 {"answers":["?"]*15},{"answers":[None]*15},{"time":"25:00"},
                 {"answers":["-"]*15},{"double":2}]
        for values in changes:
            with self.subTest(values=values),self.assertRaises((ValueError,TypeError)):
                export_report([record(**values)])

    def test_path_output_and_template_protection(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"resultado.xlsx"
            self.assertEqual(export_report([record()],path),path)
            self.assertTrue(path.is_file())
            self.assertEqual(load_workbook(path,data_only=True)["Grado 1"]["H3"].value,150)
        with self.assertRaises(ValueError):
            export_report([record()],DEFAULT_TEMPLATE)

    def test_communication_course_replaces_embedded_math_label_visually(self):
        data=export_report([record()],course="Comunicacion")
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            drawing=archive.read("xl/drawings/drawing1.xml").decode()
            self.assertIn("COMUNICACIÓN",drawing)
            self.assertIn("FFFF00",drawing)


if __name__ == "__main__":
    unittest.main()
