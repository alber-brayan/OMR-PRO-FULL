"""Prueba completa: PDF pintado -> lector real -> revisión -> Excel y PDF.

Los seis alumnos DEMO son datos artificiales y no se guardan en la sesión real.
"""
import contextlib
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest
import zipfile

import cv2
import openpyxl
import pymupdf

from app import create_app
from tests.test_omr25_engine import render_sheet, desk

ROOT = Path(__file__).resolve().parents[1]


def run_demo(output_dir=None, data_dir=None):
    """Genera seis casos artificiales a través de las mismas rutas de la app."""
    template = ROOT / "assets" / "report_template.xlsx"
    before = hashlib.sha256(template.read_bytes()).hexdigest()
    output = Path(output_dir) if output_dir else None
    if output:
        output.mkdir(parents=True, exist_ok=True)
    temp = tempfile.TemporaryDirectory() if data_dir is None else contextlib.nullcontext(str(data_dir))
    with temp as directory:
        client = create_app(directory, testing=True).test_client()
        key = "ABCDEABCDEABCDE"
        payloads = []
        expected = []
        for grade, word in enumerate(("UNO", "DOS", "TRES", "CUATRO", "CINCO", "SEIS"), 1):
            response = client.post("/api/keys", json={"course": "Matematica", "grade": grade, "answers": key})
            assert response.status_code == 200, response.json
            answers = list(key)
            if grade == 2:
                answers[0] = "B"
            elif grade == 3:
                answers[0] = "-"
            elif grade == 4:
                answers[0] = "AB"
            elif grade == 5:
                answers[0:3] = ["B", "C", "-"]
            elif grade == 6:
                answers = ["-"] * 15
            photo = desk(render_sheet(name=f"DEMO {word}", school="DEMO", grade=grade, answers=answers),
                         rotation=grade % 4, perspective=True, shadow=True)
            ok, encoded = cv2.imencode(".jpg", photo, [cv2.IMWRITE_JPEG_QUALITY, 90])
            assert ok
            payloads.append((f"DEMO_grado_{grade}.jpg", encoded.tobytes()))
            detected = ["X" if len(a) > 1 else a for a in answers]
            correct = sum(a == b for a, b in zip(detected, key))
            blank = detected.count("-")
            incorrect = 15 - correct - blank
            expected.append({"grade": grade, "name": f"DEMO {word}", "answers": detected,
                             "correct": correct, "incorrect": incorrect, "blank": blank,
                             "score": correct * 10 - incorrect})
            if output:
                (output / payloads[-1][0]).write_bytes(payloads[-1][1])

        # Un ZIP con subcarpetas imita la exportación de la cámara Android.
        zipped = io.BytesIO()
        with zipfile.ZipFile(zipped, "w", zipfile.ZIP_DEFLATED) as archive:
            for filename, raw in payloads:
                archive.writestr("Matematica/" + filename, raw)
        uploaded = client.post("/api/scan", data={"course": "Matematica",
                               "files": (io.BytesIO(zipped.getvalue()), "DEMO_seis_grados.zip")},
                               content_type="multipart/form-data")
        assert uploaded.status_code == 200, uploaded.json
        assert not uploaded.json["errors"], uploaded.json["errors"]
        records = uploaded.json["records"]
        assert len(records) == 6
        assert client.get("/api/export/excel?course=Matematica").status_code == 409
        for record in records:
            match = expected[record["grade"] - 1]
            for field in ("name", "grade", "answers", "correct", "incorrect", "blank", "score"):
                assert record[field] == match[field], (field, record, match)
            if "X" in record["answers"]:
                assert record["needs_review"], "Las dobles similares deben revisarse"
            else:
                assert not record["needs_review"], record["issues"]
            confirmed = client.post(f"/api/records/{record['id']}/confirm", json={
                "reviewed": True, "order": f"{record['grade']:03d}",
                "time": f"09:{record['grade']:02d}",
            })
            assert confirmed.status_code == 200, confirmed.json
            assert confirmed.json["confirmed"]

        duplicated = client.post("/api/scan", data={"course": "Matematica", "files":
                                  (io.BytesIO(payloads[0][1]), payloads[0][0])},
                                  content_type="multipart/form-data")
        assert duplicated.status_code == 200, duplicated.json
        assert len(client.get("/api/records?course=Matematica").json["records"]) == 6

        excel = client.get("/api/export/excel?course=Matematica")
        assert excel.status_code == 200, excel.json
        workbook = openpyxl.load_workbook(io.BytesIO(excel.data), data_only=True)
        assert len(workbook.sheetnames) == 6, workbook.sheetnames
        for grade, sheet in enumerate(workbook.worksheets, 1):
            match = expected[grade-1]
            assert sheet["B3"].value == match["name"]
            assert sheet["C3"].value == "DEMO"
            assert str(sheet["D3"].value) == f"{grade:03d}"
            for cell, field in (("E3", "correct"), ("F3", "incorrect"), ("G3", "blank"), ("H3", "score")):
                assert sheet[cell].value == match[field], (cell, sheet[cell].value, match)
            assert sheet["B4"].value is None
            for row in sheet.iter_rows():
                assert not any(c.data_type == "e" for c in row)
        workbook.close()
        pdf = client.get("/api/export/pdf?course=Matematica")
        assert pdf.status_code == 200
        with pymupdf.open(stream=pdf.data, filetype="pdf") as document:
            assert len(document) == 6
            assert document[0].rect.width > document[0].rect.height
        summary = {"synthetic": True, "description": "DEMO: seis fichas artificiales, una por grado.",
                   "cases": expected, "confirmed": 6, "duplicates_after_reupload": 0,
                   "excel_sheets": 6, "graded_pdf_pages": 6}
        if output:
            (output / "DEMO_OMR25_por_grados.xlsx").write_bytes(excel.data)
            (output / "DEMO_OMR25_calificadas.pdf").write_bytes(pdf.data)
            (output / "DEMO_validacion.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
        assert hashlib.sha256(template.read_bytes()).hexdigest() == before
        return summary


class EndToEnd25Tests(unittest.TestCase):
    def test_six_grades_real_reader_review_excel_and_pdf(self):
        result = run_demo()
        self.assertEqual(result["confirmed"], 6)


if __name__ == "__main__":
    unittest.main()
